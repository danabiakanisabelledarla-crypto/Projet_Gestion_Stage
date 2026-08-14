from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("Rapport_de_stage_InnoTechLab_FINAL.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run.add_text("Mettre à jour la table des matières dans Word (clic droit > Mettre à jour les champs).")
    run._r.append(fld_char_end)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.first_line_indent = Cm(1)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, before, after in (
        ("Title", 16, 0, 18),
        ("Heading 1", 14, 12, 10),
        ("Heading 2", 13, 10, 8),
        ("Heading 3", 12, 8, 6),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Placeholder" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Placeholder", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.bold = True
        style.font.color.rgb = RGBColor(192, 0, 0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(5)


def add_centered(doc, text, size=12, bold=False, uppercase=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text.upper() if uppercase else text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_placeholder(doc, text):
    p = doc.add_paragraph(style="Placeholder")
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run(f"[À COMPLÉTER : {text}]")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.75)
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_chapter_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1)
    p.add_run(text)


def add_chapter_conclusion(doc, text):
    doc.add_heading("Conclusion du chapitre", level=2)
    add_body(doc, text)


def build_report():
    doc = Document()
    configure_document(doc)

    # Cover
    add_centered(doc, "RÉPUBLIQUE DU CAMEROUN", 11, True)
    add_centered(doc, "Paix – Travail – Patrie", 10)
    add_centered(doc, "MINISTÈRE DE L’ENSEIGNEMENT SUPÉRIEUR", 11, True)
    add_centered(doc, "INSTITUT AFRICAIN D’INFORMATIQUE – CAMEROUN", 12, True)
    add_centered(doc, "Centre d’Excellence Technologique Paul BIYA", 11)
    doc.add_paragraph()
    add_centered(doc, "RAPPORT DE STAGE ACADÉMIQUE", 18, True)
    doc.add_paragraph()
    add_centered(doc, "SMART-HOME : CONTRÔLE DE LA CONSOMMATION D’ÉNERGIE", 16, True)
    doc.add_paragraph()
    add_centered(doc, "Stage effectué à InnoTechLab, centre d’incubation de la Digital Transformation Alliance", 12)
    add_centered(doc, "Début de l’intégration : 28 juillet 2021 – Début officiel annoncé : 2 août 2021", 11)
    add_placeholder(doc, "DATE DE FIN DU STAGE")
    doc.add_paragraph()
    add_centered(doc, "Rédigé et présenté par", 11, True)
    add_centered(doc, "TAGHE Simplice Ulrich", 14, True)
    add_centered(doc, "Étudiant en 3e année Systèmes et Réseaux", 12)
    add_placeholder(doc, "MATRICULE")
    doc.add_paragraph()
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.cell(0, 0).text = "Encadreur académique"
    table.cell(0, 1).text = "Encadreur professionnel"
    table.cell(1, 0).text = "[À COMPLÉTER : NOM ET TITRE]"
    table.cell(1, 1).text = "[À COMPLÉTER : NOM ET TITRE]"
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph()
    add_centered(doc, "Année académique 2020 – 2021", 12, True)
    page_break(doc)

    # Editorial checklist
    doc.add_heading("NOTE ÉDITORIALE À SUPPRIMER AVANT IMPRESSION", level=1)
    add_body(doc, "Le document historique fourni ne contient pas toutes les informations nécessaires pour produire un rapport techniquement complet. Afin de respecter l’exigence du guide selon laquelle seuls des faits vérifiés doivent être rapportés, les informations suivantes doivent être ajoutées par l’étudiant avant le dépôt :")
    for item in (
        "la date exacte de fin du stage ;",
        "le matricule de l’étudiant ;",
        "les noms et titres des encadreurs académique et professionnel ;",
        "le service précis d’affectation ;",
        "le journal quotidien des tâches effectivement réalisées ;",
        "les outils matériels et logiciels réellement utilisés ;",
        "le problème précis observé chez InnoTechLab ou chez l’un de ses clients ;",
        "le cahier des charges validé ;",
        "la méthode et le langage de modélisation réellement employés ;",
        "les diagrammes, choix techniques, résultats d’implémentation et tests ;",
        "les références techniques effectivement consultées ;",
        "les annexes techniques (captures, schémas, code, manuel utilisateur).",
    ):
        add_bullet(doc, item)
    add_body(doc, "Les passages en rouge et entre crochets sont des champs de travail. Ils doivent être complétés, relus et supprimés avant l’impression définitive.")
    page_break(doc)

    # Dedication
    doc.add_heading("DÉDICACE", level=1)
    add_placeholder(doc, "DÉDICACE À UNE SEULE PERSONNE OU ENTITÉ")
    page_break(doc)

    # Thanks
    doc.add_heading("REMERCIEMENTS", level=1)
    add_body(doc, "Au terme de notre stage académique, nous adressons nos remerciements aux responsables de l’Institut Africain d’Informatique du Cameroun pour la formation reçue et pour l’organisation de cette immersion professionnelle.")
    add_body(doc, "Nous exprimons également notre gratitude à la Direction Générale d’InnoTechLab, au Professeur SAMA Mbang ainsi qu’à la Direction des Ressources Humaines pour l’accueil, la présentation de la structure et l’organisation du programme de travail des stagiaires.")
    add_placeholder(doc, "REMERCIEMENTS NOMINATIFS À L’ENCADREUR ACADÉMIQUE ET AU MAÎTRE DE STAGE")
    add_body(doc, "Enfin, nous remercions le personnel et les autres stagiaires d’InnoTechLab pour les échanges intervenus durant notre intégration.")
    page_break(doc)

    # Summary
    doc.add_heading("SOMMAIRE", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_toc(p)
    page_break(doc)

    # Lists
    doc.add_heading("LISTE DES TABLEAUX", level=1)
    add_body(doc, "Tableau 1 : Fiche signalétique d’InnoTechLab")
    add_body(doc, "Tableau 2 : Ressources matérielles de la Digital Transformation Alliance")
    add_body(doc, "Tableau 3 : Journal synthétique des tâches du stage")
    page_break(doc)

    doc.add_heading("LISTE DES FIGURES", level=1)
    add_body(doc, "Figure 1 : Organigramme d’InnoTechLab – à insérer à partir du document source.")
    add_body(doc, "Figure 2 : Plan de localisation d’InnoTechLab – à insérer à partir du document source.")
    add_body(doc, "Figure 3 : Modélisation de la solution – à compléter.")
    page_break(doc)

    doc.add_heading("LISTE DES ABRÉVIATIONS", level=1)
    abbreviations = (
        ("DTA", "Digital Transformation Alliance"),
        ("IAI", "Institut Africain d’Informatique"),
        ("ONG", "Organisation Non Gouvernementale"),
        ("SABC", "Société Anonyme des Brasseries du Cameroun"),
        ("IoT", "Internet of Things, ou Internet des objets"),
    )
    for acronym, meaning in abbreviations:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(acronym).bold = True
        p.add_run(f" : {meaning}")
    page_break(doc)

    # Foreword
    doc.add_heading("AVANT-PROPOS", level=1)
    add_body(doc, "L’Institut Africain d’Informatique du Cameroun, Centre d’Excellence Technologique Paul BIYA, assure la formation de l’étudiant concerné, inscrit en troisième année de la filière Systèmes et Réseaux au cours de l’année académique 2020-2021.")
    add_body(doc, "Le stage académique complète la formation théorique par une immersion dans un environnement professionnel. Il permet à l’étudiant d’observer l’organisation d’une structure, de s’intégrer à une équipe et de confronter ses connaissances aux besoins réels d’une organisation.")
    add_placeholder(doc, "PRÉSENTATION INSTITUTIONNELLE OFFICIELLE DE L’IAI-CAMEROUN")
    page_break(doc)

    # Introduction
    doc.add_heading("INTRODUCTION GÉNÉRALE", level=1)
    add_body(doc, "La transformation numérique touche progressivement les domaines de l’éducation, de l’industrie, de l’agriculture, de la santé et de l’énergie. Dans ce contexte, les structures d’innovation jouent un rôle important. Elles offrent des espaces de formation, de prototypage et d’accompagnement de projets. InnoTechLab s’inscrit dans cette dynamique au Cameroun.")
    add_body(doc, "Le thème associé au stage est « Smart-Home : contrôle de la consommation d’énergie ». Une maison intelligente désigne un habitat dans lequel des équipements peuvent être surveillés ou commandés au moyen de dispositifs numériques. Le contrôle de la consommation d’énergie renvoie à l’observation, à la mesure et à la maîtrise de l’énergie utilisée par ces équipements.")
    add_body(doc, "Le document historique fourni présente l’accueil du stagiaire et la structure d’InnoTechLab. Il ne décrit cependant ni le problème technique détaillé, ni le cahier des charges, ni l’implémentation d’une solution. La problématique définitive doit donc être formulée à partir des faits réellement observés pendant le stage.")
    add_placeholder(doc, "PROBLÉMATIQUE VALIDÉE À PARTIR DU PROBLÈME RÉELLEMENT OBSERVÉ")
    add_body(doc, "L’intérêt de l’étude doit être apprécié à deux niveaux. Sur le plan pédagogique, elle doit permettre la mobilisation des connaissances acquises en Systèmes et Réseaux. Sur le plan professionnel, elle doit contribuer à la compréhension ou à la résolution d’un besoin identifié au sein de la structure d’accueil ou chez l’un de ses clients.")
    add_placeholder(doc, "OBJECTIF GÉNÉRAL ET OBJECTIFS SPÉCIFIQUES VALIDÉS")
    add_body(doc, "Le présent rapport est structuré en trois chapitres. Le premier présente InnoTechLab et le déroulement du stage. Le deuxième identifie le problème étudié et examine les solutions existantes. Le troisième expose la solution effectivement proposée, sa modélisation, sa réalisation et ses résultats.")
    page_break(doc)

    # Chapter 1
    doc.add_heading("CHAPITRE 1 : PRÉSENTATION DE L’ENTREPRISE ET DÉROULEMENT DU STAGE", level=1)
    add_chapter_intro(doc, "Ce chapitre présente d’abord InnoTechLab dans son environnement interne et externe. Il décrit ensuite l’accueil, l’intégration et les activités du stage à partir des informations disponibles.")

    doc.add_heading("Section 1 : Présentation d’InnoTechLab", level=2)
    doc.add_heading("1.1. Fiche signalétique", level=3)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = table.rows[0].cells
    headers[0].text = "Rubrique"
    headers[1].text = "Information disponible"
    set_repeat_table_header(table.rows[0])
    for cell in headers:
        set_cell_shading(cell, "D9EAF7")
    data = (
        ("Nom", "InnoTechLab"),
        ("Nature", "Centre camerounais d’incubation et d’innovation technologique, pédagogique et industrielle"),
        ("Date d’inauguration", "14 octobre 2020"),
        ("Organisation de soutien", "Digital Transformation Alliance, organisation internationale à but non lucratif"),
        ("Direction mentionnée", "Professeur SAMA Mbang, Directeur Général"),
        ("Siège", "Mvan, carrefour SABC, Yaoundé"),
        ("Domaines mentionnés", "Éducation, transformation industrielle, agriculture, santé et énergie"),
        ("Forme juridique et coordonnées complètes", "[À COMPLÉTER À PARTIR DE DOCUMENTS OFFICIELS]"),
    )
    for label, value in data:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    caption = doc.add_paragraph("Tableau 1 : Fiche signalétique d’InnoTechLab")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    source = doc.add_paragraph("Source : document « Historique InnoTechLab » fourni par l’étudiant.")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.first_line_indent = Cm(0)

    doc.add_heading("1.2. Historique et missions", level=3)
    add_body(doc, "InnoTechLab a été inauguré le 14 octobre 2020. Le centre est dédié à l’éducation et à la transformation industrielle. Il est soutenu par la Digital Transformation Alliance, organisation internationale à but non lucratif dirigée par le Professeur SAMA Mbang.")
    add_body(doc, "Le centre constitue un espace de prototypage, d’innovation et de démonstration. Il vise la maîtrise des technologies et de l’ingénierie. Ses activités concernent notamment l’agriculture, la santé et l’énergie. Il sert aussi d’espace d’incubation, d’éducation et de formation. À ce titre, il peut accueillir des événements, des conférences et des camps de formation.")
    add_body(doc, "Selon le document historique, InnoTechLab s’adresse aux apprenants, aux chercheurs, aux entrepreneurs et aux chefs d’entreprise. Les apprenants peuvent y accéder à des formations. Les chercheurs peuvent y rechercher des solutions à des problèmes rencontrés en Afrique. Les entrepreneurs peuvent y développer leurs projets. Les responsables d’entreprise peuvent y échanger avec une communauté d’innovateurs et tester des idées.")

    doc.add_heading("1.3. Organisation et fonctionnement", level=3)
    add_body(doc, "L’organisation décrite comprend un Conseil d’Administration, une Direction Générale, un Secrétariat Général, une Direction Pédagogique et de la Formation, un responsable de l’industrialisation ainsi qu’une Direction de la Communication et du Marketing.")
    add_body(doc, "Le Conseil d’Administration définit les orientations stratégiques, traite les questions liées au fonctionnement et exerce une mission de contrôle. Le Directeur Général représente la structure et exerce les pouvoirs qui lui sont confiés. Le Secrétariat Général assure notamment la gestion des courriers, des agendas, des réunions et du classement documentaire.")
    add_body(doc, "La Direction Pédagogique et de la Formation traduit les orientations générales en stratégie de formation. Le responsable de l’industrialisation planifie les méthodes techniques et les moyens nécessaires à la fabrication de produits. La Direction de la Communication et du Marketing conçoit et pilote la stratégie de communication et de marketing.")
    add_placeholder(doc, "INSÉRER LA FIGURE 1 : ORGANIGRAMME D’INNOTECHLAB, AVEC TITRE ET SOURCE SOUS LA FIGURE")

    doc.add_heading("1.4. Ressources matérielles et logicielles", level=3)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("Équipement", "Marque", "Quantité", "Caractéristiques")):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), "D9EAF7")
    set_repeat_table_header(table.rows[0])
    equipment = (
        ("Ordinateurs de bureau", "Acer, Asus", "11", "Core i5, 4e génération, 32 Go de mémoire vive, stockage 1 To, carte graphique Nvidia dédiée 2 Go"),
        ("Ordinateurs portables", "HP", "10", "Core i5, 11e génération, 32 Go de mémoire vive, stockage 500 Go"),
        ("Ordinateurs portables", "HP", "5", "Core i7, 5e génération, 32 Go de mémoire vive, stockage 500 Go"),
        ("Routeurs", "Fritz!Box, Cisco", "2", "Caractéristiques non précisées dans le document source"),
    )
    for row_values in equipment:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
    caption = doc.add_paragraph("Tableau 2 : Ressources matérielles de la Digital Transformation Alliance")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    source = doc.add_paragraph("Source : document « Historique InnoTechLab » fourni par l’étudiant.")
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.first_line_indent = Cm(0)
    add_body(doc, "Les logiciels cités dans le document historique sont SEE Electrical, PVsyst et Solarius. Le document ne précise pas les versions, les licences ni les usages effectués par le stagiaire.")

    doc.add_heading("1.5. Partenaires mentionnés", level=3)
    add_body(doc, "Les partenaires académiques locaux mentionnés sont l’Institut Africain d’Informatique du Cameroun, IFTIC-SUP, UPAC et JFN. Les partenaires académiques internationaux cités sont l’Université de Lorraine, l’Université Polytechnique de Nancy et l’ESSTI.")
    add_body(doc, "Les partenaires industriels et les entreprises mentionnés sont ALUCAM, CICAM, SARMETAL et Digit-Tech-Innov Solutions. Le document cite également Dassault Systèmes et Microsoft Nonprofit parmi les partenaires relevant des organisations non gouvernementales et des fondations.")

    doc.add_heading("Section 2 : Accueil, intégration et activités du stage", level=2)
    doc.add_heading("2.1. Accueil et intégration", level=3)
    add_body(doc, "Notre intégration au sein d’InnoTechLab a commencé le 28 juillet 2021. La Directrice des Ressources Humaines nous a souhaité la bienvenue avant de nous conduire dans la salle de conférence.")
    add_body(doc, "Une réunion a ensuite été organisée en visioconférence sous la conduite du Professeur SAMA Mbang, Directeur Général d’InnoTechLab. Cette réunion a permis de présenter la structure et son personnel. Elle s’est achevée par l’établissement du programme de travail et l’affectation des maîtres de stage. Le document historique indique que le 2 août 2021 a marqué le début officiel du stage.")

    doc.add_heading("2.2. Service d’affectation et tâches réalisées", level=3)
    add_body(doc, "Le document historique ne précise pas le service d’affectation, le nom du maître de stage, les tâches quotidiennes, les outils utilisés ni les résultats obtenus. Ces éléments doivent être fournis à partir du cahier de stage. Ils ne peuvent pas être déduits de la seule présentation d’InnoTechLab.")
    add_placeholder(doc, "SERVICE OU DIRECTION D’AFFECTATION ET POSITION DANS L’ORGANIGRAMME")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("Date / jour", "Lieu", "Tâche effectuée", "Outils utilisés", "Personnes impliquées")):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for _ in range(8):
        cells = table.add_row().cells
        for cell in cells:
            cell.text = "[À COMPLÉTER]"
    caption = doc.add_paragraph("Tableau 3 : Journal synthétique des tâches du stage")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)

    doc.add_heading("2.3. Problème observé", level=3)
    add_body(doc, "Le thème du stage porte sur le contrôle de la consommation d’énergie dans une maison intelligente. Toutefois, le document source ne présente pas le besoin du client, les causes du problème, ses manifestations ni son impact sur InnoTechLab. Ces éléments doivent être décrits à partir des observations et des échanges réellement intervenus pendant le stage.")
    add_placeholder(doc, "DESCRIPTION FACTUELLE DU PROBLÈME, DE SES CAUSES, DE SES MANIFESTATIONS ET DE SON IMPACT")
    add_chapter_conclusion(doc, "Ce chapitre a présenté les informations vérifiables relatives à InnoTechLab et à l’intégration du stagiaire. Il a aussi mis en évidence les informations du déroulement technique qui doivent être complétées à partir du cahier de stage.")
    page_break(doc)

    # Chapter 2
    doc.add_heading("CHAPITRE 2 : IDENTIFICATION DU PROBLÈME", level=1)
    add_chapter_intro(doc, "Ce chapitre doit préciser le besoin réel de l’entreprise ou de l’un de ses clients. Il doit ensuite présenter les solutions existantes et leurs limites par rapport au problème étudié. Les données techniques nécessaires ne figurent pas dans le document historique. La structure suivante est donc prête à être complétée avec des faits vérifiés.")

    doc.add_heading("Section 1 : Analyse du besoin", level=2)
    doc.add_heading("1.1. Contexte spécifique", level=3)
    add_placeholder(doc, "CONTEXTE DU CLIENT, SITE CONCERNÉ, UTILISATEURS ET SITUATION OBSERVÉE")
    doc.add_heading("1.2. Description du problème", level=3)
    add_placeholder(doc, "PROBLÈME PRÉCIS LIÉ AU CONTRÔLE DE LA CONSOMMATION D’ÉNERGIE")
    doc.add_heading("1.3. Causes, manifestations et impacts", level=3)
    add_placeholder(doc, "CAUSES VÉRIFIÉES")
    add_placeholder(doc, "MANIFESTATIONS MESURÉES OU OBSERVÉES")
    add_placeholder(doc, "IMPACT SUR L’ENTREPRISE OU LE CLIENT")
    doc.add_heading("1.4. Expression du besoin", level=3)
    add_placeholder(doc, "BESOIN FONCTIONNEL, UTILISATEURS, CONTRAINTES ET CRITÈRES D’ACCEPTATION")

    doc.add_heading("Section 2 : Étude des solutions existantes", level=2)
    doc.add_heading("2.1. Solutions identifiées", level=3)
    add_placeholder(doc, "SOLUTIONS RÉELLEMENT ÉTUDIÉES, AVEC SOURCES CITÉES")
    doc.add_heading("2.2. Analyse comparative", level=3)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("Solution", "Fonctions", "Avantages", "Limites", "Écart avec le besoin")):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for _ in range(4):
        cells = table.add_row().cells
        for cell in cells:
            cell.text = "[À COMPLÉTER]"
    doc.add_heading("2.3. Valeur ajoutée attendue", level=3)
    add_placeholder(doc, "VALEUR AJOUTÉE DE LA SOLUTION PROPOSÉE PAR RAPPORT AUX SOLUTIONS EXISTANTES")
    add_chapter_conclusion(doc, "L’identification du problème doit conduire à un cahier des charges précis. La version définitive de ce chapitre devra être fondée sur des observations, des mesures, des entretiens ou des documents techniques cités.")
    page_break(doc)

    # Chapter 3
    doc.add_heading("CHAPITRE 3 : SOLUTION PROPOSÉE", level=1)
    add_chapter_intro(doc, "Ce chapitre doit présenter l’apport pratique du stagiaire. Conformément au guide, il doit décrire la démarche de modélisation, le langage de modélisation, les outils utilisés, la conception, l’implémentation et les résultats. Aucune de ces informations n’étant détaillée dans le document historique, elles doivent être renseignées à partir du travail réellement effectué.")

    doc.add_heading("Section 1 : Analyse et conception de la solution", level=2)
    doc.add_heading("1.1. Cahier des charges", level=3)
    add_placeholder(doc, "OBJECTIFS, PÉRIMÈTRE, ACTEURS, EXIGENCES FONCTIONNELLES ET NON FONCTIONNELLES")
    doc.add_heading("1.2. Démarche et langage de modélisation", level=3)
    add_placeholder(doc, "MÉTHODE DE MODÉLISATION RÉELLEMENT UTILISÉE")
    add_placeholder(doc, "LANGAGE ET DIAGRAMMES RÉELLEMENT PRODUITS")
    doc.add_heading("1.3. Architecture retenue", level=3)
    add_placeholder(doc, "SCHÉMA D’ARCHITECTURE MATÉRIELLE, LOGICIELLE ET RÉSEAU")
    doc.add_heading("1.4. Choix des outils et technologies", level=3)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("Outil / technologie", "Rôle", "Version", "Justification du choix")):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for _ in range(6):
        cells = table.add_row().cells
        for cell in cells:
            cell.text = "[À COMPLÉTER]"

    doc.add_heading("Section 2 : Implémentation, tests et exploitation", level=2)
    doc.add_heading("2.1. Réalisation de la solution", level=3)
    add_placeholder(doc, "ÉTAPES D’IMPLÉMENTATION, CONFIGURATIONS, CODE ET CAPTURES COMMENTÉES")
    doc.add_heading("2.2. Scénarios et résultats de test", level=3)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("Test", "Précondition", "Action", "Résultat attendu", "Résultat obtenu")):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), "D9EAF7")
    set_repeat_table_header(table.rows[0])
    for _ in range(6):
        cells = table.add_row().cells
        for cell in cells:
            cell.text = "[À COMPLÉTER]"
    doc.add_heading("2.3. Manuel utilisateur", level=3)
    add_placeholder(doc, "PROCÉDURE D’INSTALLATION, DE CONNEXION, D’UTILISATION ET DE DÉPANNAGE")
    doc.add_heading("2.4. Limites et perspectives techniques", level=3)
    add_placeholder(doc, "LIMITES CONSTATÉES LORS DES TESTS ET AMÉLIORATIONS ENVISAGÉES")
    add_chapter_conclusion(doc, "La version finale de ce chapitre devra démontrer la contribution effective du stagiaire au moyen de schémas, de tableaux de tests, de captures et d’explications techniques vérifiables.")
    page_break(doc)

    # Conclusion
    doc.add_heading("CONCLUSION GÉNÉRALE", level=1)
    add_body(doc, "Le stage effectué au sein d’InnoTechLab a commencé par une phase d’accueil et d’intégration. Cette phase a permis de découvrir un centre camerounais d’incubation et d’innovation dédié notamment à l’éducation et à la transformation industrielle. Elle a aussi permis de prendre connaissance de son organisation, de ses missions, de ses ressources et de certains de ses partenaires.")
    add_body(doc, "Le thème du stage, « Smart-Home : contrôle de la consommation d’énergie », inscrit le travail dans le domaine des systèmes numériques appliqués à l’énergie. Toutefois, les documents fournis ne décrivent pas les tâches techniques, la démarche de modélisation, l’implémentation ni les résultats. Le récapitulatif technique et les contributions du stagiaire devront donc être complétés à partir du cahier de stage et des livrables réellement produits.")
    add_placeholder(doc, "BREF RÉCAPITULATIF DES RÉSULTATS OBTENUS ET CONTRIBUTIONS PERSONNELLES")
    add_placeholder(doc, "DÉMARCHE EFFECTIVEMENT SUIVIE")
    add_placeholder(doc, "PERSPECTIVES D’AMÉLIORATION DU TRAVAIL")
    page_break(doc)

    # References
    doc.add_heading("RÉFÉRENCES BIBLIOGRAPHIQUES", level=1)
    add_body(doc, "TAGHE, S. U. (2021), « Historique InnoTechLab », document de stage non publié, Yaoundé.")
    add_body(doc, "INSTITUT SAINT JEAN (2024), « Guide de rédaction du rapport de stage projet », document pédagogique, Yaoundé.")
    add_placeholder(doc, "OUVRAGES, ARTICLES ET SOURCES ÉLECTRONIQUES EFFECTIVEMENT CITÉS DANS LES CHAPITRES 2 ET 3")
    page_break(doc)

    # Annexes
    doc.add_heading("ANNEXES", level=1)
    doc.add_heading("Annexe A : Organigramme d’InnoTechLab", level=2)
    add_placeholder(doc, "INSÉRER L’ORGANIGRAMME DU DOCUMENT SOURCE")
    page_break(doc)
    doc.add_heading("Annexe B : Plan de localisation d’InnoTechLab", level=2)
    add_placeholder(doc, "INSÉRER LE PLAN DE LOCALISATION DU DOCUMENT SOURCE")
    page_break(doc)
    doc.add_heading("Annexe C : Pièces techniques du projet", level=2)
    add_placeholder(doc, "INSÉRER LES DIAGRAMMES, SCHÉMAS, CAPTURES, EXTRAITS DE CODE ET RÉSULTATS DE TEST")
    page_break(doc)

    # Full TOC
    doc.add_heading("TABLE DES MATIÈRES", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_toc(p)
    page_break(doc)

    # Evaluation sheet, requested after page 17
    doc.add_heading("FICHE D’ÉVALUATION DE STAGE", level=1)
    info = doc.add_table(rows=4, cols=2)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = (
        ("Année académique", "2020-2021"),
        ("Cycle / niveau", "3e année"),
        ("Filière / option", "Systèmes et Réseaux"),
        ("Candidat et intitulé", "TAGHE Simplice Ulrich – Smart-Home : contrôle de la consommation d’énergie"),
    )
    for row, (label, value) in zip(info.rows, labels):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "D9EAF7")
    doc.add_paragraph()
    evaluation = doc.add_table(rows=1, cols=3)
    evaluation.style = "Table Grid"
    evaluation.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(("A – Rapport de stage [45 %]", "Barème", "Note")):
        evaluation.cell(0, i).text = value
        set_cell_shading(evaluation.cell(0, i), "D9EAF7")
    criteria = (
        ("Fond – Compréhension du besoin du client", "/ 2"),
        ("Fond – Analyse / Conception / Implémentation", "/ 5"),
        ("Fond – Utilisation des connaissances scientifiques et techniques", "/ 4"),
        ("Fond – Exploitation du travail effectué", "/ 3"),
        ("Forme – Présentation et clarté", "/ 3"),
        ("Forme – Manuel utilisateur / Rapport technique / Bibliographie", "/ 3"),
        ("Total rubrique Rapport", "/ 20"),
        ("Discipline – Pénalité de retard de dépôt", ""),
        ("Note finale du rapport", "/ 20"),
        ("B – Note de soutenance [35 %]", "/ 20"),
        ("C – Note d’entreprise [20 %]", "/ 20"),
        ("Note finale du stage", "/ 20"),
    )
    for criterion, scale in criteria:
        cells = evaluation.add_row().cells
        cells[0].text = criterion
        cells[1].text = scale
        cells[2].text = ""
    doc.add_paragraph()
    signatures = doc.add_table(rows=2, cols=3)
    signatures.style = "Table Grid"
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(("Responsable pédagogique", "Président du jury", "Directeur de l’établissement")):
        signatures.cell(0, index).text = value
        signatures.cell(1, index).text = "\n\nSignature :"
        set_cell_shading(signatures.cell(0, index), "D9EAF7")
        for p in signatures.cell(0, index).paragraphs + signatures.cell(1, index).paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)

    # Document field setting: update fields on open.
    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_report()
